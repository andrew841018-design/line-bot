"""美股盤勢的資料來源（/盤勢，2026-10-17）：Nasdaq、SSGA、MoneyDJ、CBOE、美國財政部、CFTC、FINRA、Fed。

Andrew 2026-10-17：「美股/台股的籌碼/基本面/技術面/消息面/期權市場的下注（最大未平倉價位等等，gamma
squeeze…)/持股跌破20MA,60MA,120MA,240MA的家數/成交量…我要知道當前量是多還是少/持股上漲家數與下跌家數
比例/13F相關資訊/市場上槓桿的程度…」，標的 0050／SOXX／VOO／QQQ。這個模組只負責美股那一半的「抓」與
「讀懂」：計算在 market_calc、組卡在 market_dashboard、排程與存檔在 jobs/market_snapshot.py。

每個來源兩個函式（INTERFACES §3）：
- 純解析 `parse_*`：收 bytes／str／已解析的 JSON，回 dataclass；格式不對丟 `ParseError(kind)`，kind 是
  format、missing_field、empty、too_few、date、conflict 之一（不帶回應內容）。
- 抓取 `fetch_*`：組網址 → `http.fetch(url, max_bytes=…, timeout=…, accept=(…), ledger=ledger)`（預設
  market_http，測試注入假的）→ 解析。每個端點的大小、時間、Content-Type、每日預算在 `ENDPOINTS`。

口徑（PLAN §3.1、§18、§19）：
- 代號一律 canonical：美股大寫、股別用點（BRK.B）。SSGA、CBOE 原樣；Nasdaq 原樣（篩選器那類端點寫
  「BRK/B」，一併換成點）；MoneyDJ 去「.US」；Yahoo 的「BRK-B」換成點。台股 4–6 碼數字（MoneyDJ 去「.TW」）。
- 持股表的非股票列（代號不合格、SSGA 識別碼不是有效 CUSIP、名稱有現金／期貨／交換／擔保品／貨幣市場）
  濾掉並記數；MoneyDJ 沒有代號連結的列查固定的名稱→代號表（ADR 對到美國代號），對不到的記檔數與權重。
- 時間序列一律依日期由舊到新；同一個鍵重複且值不同 → ParseError("conflict")。
- 時區：CBOE 頂層 timestamp 是 UTC、last_trade_time 是美東當地時間（沒帶時區）。Nasdaq 經濟日曆的
  `gmt` 欄其實是美東當地時間（實測 fixture：CPI 08:30、褐皮書 14:00、澳洲就業 20:30＝雪梨 11:30）。

安全（PLAN §19.2）：網址裡的代號先驗 canonical 格式再 quote(safe='')；xlsx 先檢查 zip 目錄（成員 ≤64、
解壓總量 ≤50 MB），XML 成員不得含 DOCTYPE／ENTITY、不得是 UTF-16，才交給 openpyxl（read_only、
data_only）。持股名稱（Holding.name）、CFTC 市場名這裡只做正規化（NFKC、去控制字元、壓空白），進 SQLite
前與組卡前由 job／bot 再過 `market_calc.clean_external`（PLAN §19.1）；對不到代號的名稱（unmapped_names）
會整份存進 SQLite，所以這裡就先清洗。

不 import main、stock_picks、fcn_eval、card_pages；market_http、openpyxl 用到才載入；import 時不做任何 I/O。
"""

from __future__ import annotations

import csv
import functools
import io
import json
import math
import re
import unicodedata
import warnings
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from html.parser import HTMLParser
from typing import Any, Iterable, Iterator, Optional
from urllib.parse import parse_qs, quote, urlencode, urlsplit
from zoneinfo import ZoneInfo

import market_calc
from market_calc import ParseError  # 全專案共用（2026-10-18 拿掉「還沒有 market_calc」時的本地版本）


UTC = timezone.utc
KIB = 1024
MIB = 1024 * 1024


def _err(kind: str, detail: str = "") -> ParseError:
    """只用一個位置參數建立 ParseError（共用版本的建構子可能只收 kind）；細節只放欄位名這類固定字。"""
    error = ParseError(kind)
    if detail:
        try:
            error.detail = detail
        except AttributeError:
            pass
    return error


@functools.cache
def _ny() -> ZoneInfo:
    """美東時區：第一次用才讀時區檔，import 時不做 I/O。"""
    return ZoneInfo("America/New_York")


# ── 網址與端點 ────────────────────────────────────────────────────────────
NASDAQ_API = "https://api.nasdaq.com/api"
# 主 Claude 實測：/us/en/intermediary/… 的舊路徑會 301，market_http 不轉址，所以寫死最終網址。
SSGA_SPY_URL = "https://www.ssga.com/library-content/products/fund-data/etfs/us/holdings-daily-us-en-spy.xlsx"
MONEYDJ_URL = "https://www.moneydj.com/ETF/X/Basic/Basic0007B.xdjhtm"  # 「全部持股」頁
# cdn.cboe.com 會 307 到 cdn-api（PLAN v1 實測），直接打最終主機。
CBOE_OPTIONS_URL = "https://cdn-api.cboe.com/api/global/delayed_quotes/options/{symbol}.json"
CBOE_QUOTE_URL = "https://cdn-api.cboe.com/api/global/delayed_quotes/quotes/_{symbol}.json"
TREASURY_URL = (
    "https://home.treasury.gov/resource-center/data-chart-center/interest-rates/daily-treasury-rates.csv/all/{ym}"
    "?type=daily_treasury_yield_curve&field_tdr_date_value_month={ym}&page&_format=csv"
)
CFTC_TFF_URL = "https://www.cftc.gov/dea/newcot/FinFutWk.txt"
FINRA_MARGIN_URL = "https://www.finra.org/sites/default/files/2021-03/margin-statistics.xlsx"
FED_FOMC_URL = "https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm"

CBOE_OPTION_SYMBOLS = ("QQQ", "SPY", "SOXX")
CBOE_QUOTE_SYMBOLS = ("VIX", "VIX3M")


@dataclass(frozen=True)
class Endpoint:
    """一個端點的抓取設定與每日預算（PLAN §9：估計的起始值；job 依此排程，帳本另外依主機預扣）。"""

    host: str
    accept: tuple[str, ...]  # Content-Type 要含的子字串（主 Claude 10/17 實測的型別）
    max_bytes: int
    timeout: float  # 整個請求的總時間上限（秒）：一般 15 秒、CBOE／SSGA 30 秒（PLAN §19.2）
    per_day: int  # 正常每天最多幾次
    bootstrap_per_day: int  # bootstrap 期間每天最多幾次
    retries: int  # 同一個目標（同一檔、同一天）失敗後最多再排幾次
    per_run: Optional[int] = None  # 每一輪 collect 最多幾次（沒有就是不另外限制）


ENDPOINTS: dict[str, Endpoint] = {
    # 成分聯集約 520 檔＋VOO／QQQ／SOXX／SPY＋宏觀代理 8 檔；沒有 D 列的 30 分鐘後重排、最多 4 次（§18.5）
    "nasdaq_historical": Endpoint("api.nasdaq.com", ("json",), 512 * KIB, 15.0, 560, 600, 4, per_run=60),
    "nasdaq_market_info": Endpoint("api.nasdaq.com", ("json",), 64 * KIB, 15.0, 4, 4, 2),
    "nasdaq_splits": Endpoint("api.nasdaq.com", ("json",), 2 * MIB, 15.0, 1, 2, 2),
    "nasdaq_nasdaq100": Endpoint("api.nasdaq.com", ("json",), 512 * KIB, 15.0, 1, 1, 2),
    "nasdaq_13f": Endpoint("api.nasdaq.com", ("json",), 256 * KIB, 15.0, 30, 30, 1),  # 三檔前十大合併 ≤30
    "nasdaq_financials": Endpoint("api.nasdaq.com", ("json",), 512 * KIB, 15.0, 30, 30, 1),
    # 前十大裡不在 Nasdaq-100 名單的公司市值（BRK.B、JPM、LLY、TSM…）：每週一次，三檔合併 ≤30 家；回應約 1.1 KB
    "nasdaq_summary": Endpoint("api.nasdaq.com", ("json",), 64 * KIB, 15.0, 30, 30, 1),
    "nasdaq_econ_calendar": Endpoint("api.nasdaq.com", ("json",), 1 * MIB, 15.0, 7, 7, 1),  # 7 天各一次
    "nasdaq_earnings": Endpoint("api.nasdaq.com", ("json",), 64 * KIB, 15.0, 30, 30, 1),
    "ssga_spy": Endpoint("www.ssga.com", ("spreadsheet",), 4 * MIB, 30.0, 2, 2, 1),
    "moneydj": Endpoint("www.moneydj.com", ("html",), 2 * MIB, 15.0, 4, 4, 1),  # QQQ、SOXX、VOO、0050
    "cboe_options": Endpoint("cdn-api.cboe.com", ("json",), 16 * MIB, 30.0, 3, 3, 1),  # SPY 實測 5.8 MB
    "cboe_quote": Endpoint("cdn-api.cboe.com", ("json",), 64 * KIB, 15.0, 4, 4, 1),  # VIX、VIX3M
    "treasury": Endpoint("home.treasury.gov", ("csv",), 1 * MIB, 15.0, 2, 2, 2),  # 本月＋月初時的上個月
    "cftc_tff": Endpoint("www.cftc.gov", ("text/plain",), 2 * MIB, 15.0, 1, 1, 2),
    "finra_margin": Endpoint("www.finra.org", ("spreadsheet",), 4 * MIB, 15.0, 1, 1, 2),
    "fed_fomc": Endpoint("www.federalreserve.gov", ("html",), 2 * MIB, 15.0, 1, 1, 2),
}


# ── 代號 ──────────────────────────────────────────────────────────────────
US_SYMBOL_RE = re.compile(r"^[A-Z]{1,5}(\.[A-Z])?$")
TW_CODE_RE = re.compile(r"^[0-9]{4,6}$")
OPTION_SYMBOL_RE = re.compile(r"^(QQQ|SPY|SOXX)(\d{6})([CP])(\d{8})$")
_CUSIP_RE = re.compile(r"^[0-9A-Z*@#]{8}[0-9]$")


def canonical_us_symbol(raw: object, source: str = "nasdaq") -> Optional[str]:
    """各來源的美股代號 → canonical（大寫、股別用點）；不合格回 None。

    source：ssga、nasdaq、cboe（原樣；「/」換成點）、moneydj（要有「.US」並去掉）、yahoo（「-」換成點）。"""
    if not isinstance(raw, str):
        return None
    text = unicodedata.normalize("NFKC", raw).strip().upper()
    if source == "moneydj":
        if not text.endswith(".US"):
            return None
        text = text[:-3]
    elif source == "yahoo":
        text = text.replace("-", ".")
    elif source in ("ssga", "nasdaq", "cboe"):
        text = text.replace("/", ".")
    else:
        raise ValueError(f"unknown symbol source: {source!r}")
    return text if US_SYMBOL_RE.fullmatch(text) else None


def to_source_symbol(symbol: str, source: str) -> str:
    """canonical → 各來源的寫法（Yahoo 用「-」、MoneyDJ 加「.US」、其他原樣）。"""
    if not isinstance(symbol, str) or not US_SYMBOL_RE.fullmatch(symbol):
        raise ValueError("not a canonical US symbol")
    if source == "yahoo":
        return symbol.replace(".", "-")
    if source == "moneydj":
        return symbol + ".US"
    if source in ("ssga", "nasdaq", "cboe"):
        return symbol
    raise ValueError(f"unknown symbol source: {source!r}")


def _path_symbol(symbol: object) -> str:
    """網址路徑裡的代號：代號來自第三方名單，先驗 canonical 再 quote(safe='')（PLAN §19.2）。"""
    if not isinstance(symbol, str) or not US_SYMBOL_RE.fullmatch(symbol):
        raise _err("format", "symbol")
    return quote(symbol, safe="")


# ── 文字、數字、日期 ──────────────────────────────────────────────────────
_DROP_CATEGORIES = frozenset({"Cc", "Cf", "Cs", "Co", "Cn", "Zl", "Zp"})
_MISSING_TOKENS = frozenset({"", "-", "--", "---", "N/A", "NA", "NAN", "NONE", "NULL", "&NBSP;"})
_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_MDY_RE = re.compile(r"^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$")
_MON_D_Y_RE = re.compile(r"([A-Za-z]{3})[A-Za-z]*\.?\s+(\d{1,2}),\s*(\d{4})")


def _norm_text(value: object, max_len: int = 80) -> str:
    """外部名稱的正規化（不是安全清洗）：NFKC、去控制／格式／私用字元、壓空白、截長度。"""
    if not isinstance(value, str):
        return ""
    text = unicodedata.normalize("NFKC", value)
    text = "".join(
        " " if ch in "\t\r\n" else ch
        for ch in text
        if ch in "\t\r\n" or unicodedata.category(ch) not in _DROP_CATEGORIES
    )
    return " ".join(text.split())[:max_len]


def _match_key(name: str) -> str:
    """名稱比對用：大寫、英數字與中文以外都當空白（「Brown-Forman」＝「BROWN FORMAN」）。"""
    upper = _norm_text(name, 200).upper()
    return " ".join(re.sub(r"[^0-9A-Z一-鿿]+", " ", upper).split())


def _num(value: object) -> Optional[float]:
    """來源的數字 → float。缺值（None、空白、「--」「N/A」、NaN）回 None；看不懂的字丟 format。

    收「$343.46」「59,946,370」「+0.75%」「($1,234)」「-$1,234」「$-1,234」。"""
    if value is None:
        return None
    if isinstance(value, bool):
        raise _err("format", "number")
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    if not isinstance(value, str):
        raise _err("format", "number")
    text = value.strip().replace("\xa0", "")
    if text.upper() in _MISSING_TOKENS:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1].strip()
        if text.startswith(("-", "+")):
            raise _err("format", "number")  # 「(-1,234)」兩種負號一起出現：不猜
    text = text.replace(",", "").replace("$", "").replace("%", "").strip()
    if not re.fullmatch(r"[+-]?(\d+(\.\d*)?|\.\d+)([eE][+-]?\d+)?", text):
        raise _err("format", "number")
    number = float(text)
    if not math.isfinite(number):
        return None
    return -number if negative else number


def _num_or_none(value: object) -> Optional[float]:
    """次要欄位：看不懂就當缺值，不讓一個怪欄位擋掉整份資料。"""
    try:
        return _num(value)
    except ParseError:
        return None


def _int(value: object) -> Optional[int]:
    """整數（量、口數、家數）：小數部分必須是 0。"""
    number = _num(value)
    if number is None:
        return None
    if number != math.floor(number):
        raise _err("format", "integer")
    return int(number)


def _count_or_none(value: object) -> Optional[int]:
    """次要的非負整數（量）：看不懂或是負數就當缺值。"""
    try:
        number = _int(value)
    except ParseError:
        return None
    return number if number is not None and number >= 0 else None


def _positive_or_none(value: object) -> Optional[float]:
    number = _num_or_none(value)
    return number if number is not None and number > 0 else None


def _safe_date(year: int, month: int, day: int, field_name: str) -> date:
    if not 1900 <= year <= 2200:
        raise _err("date", field_name)
    try:
        return date(year, month, day)
    except ValueError:
        raise _err("date", field_name) from None


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _mdy(value: object, field_name: str) -> date:
    """「10/16/2026」「9/30/2026」→ date；空的丟 missing_field、格式不對丟 date。"""
    if _is_blank(value):
        raise _err("missing_field", field_name)
    match = _MDY_RE.fullmatch(value) if isinstance(value, str) else None
    if not match:
        raise _err("date", field_name)
    return _safe_date(int(match.group(3)), int(match.group(1)), int(match.group(2)), field_name)


def _mon_d_y(value: object, field_name: str) -> date:
    """「Oct 16, 2026」→ date（月份自己查表，不靠 locale）。"""
    if _is_blank(value):
        raise _err("missing_field", field_name)
    match = _MON_D_Y_RE.fullmatch(value.strip()) if isinstance(value, str) else None
    month = _MONTHS.get(match.group(1).lower()) if match else None
    if not match or month is None:
        raise _err("date", field_name)
    return _safe_date(int(match.group(3)), month, int(match.group(2)), field_name)


def _aware(value: object, zone: Any, field_name: str) -> datetime:
    """ISO 時間（「2026-10-16 16:33:11」「2026-10-16T12:18:10」）；沒帶時區就當 `zone` 的當地時間。"""
    if _is_blank(value):
        raise _err("missing_field", field_name)
    if not isinstance(value, str):
        raise _err("date", field_name)
    try:
        parsed = datetime.fromisoformat(value.strip())
    except ValueError:
        raise _err("date", field_name) from None
    if not 1990 <= parsed.year <= 2200:
        raise _err("date", field_name)
    if parsed.tzinfo is None:
        return parsed.replace(tzinfo=zone)
    return parsed.astimezone(zone)


def _as_text(value: object) -> str:
    """bytes 用 UTF-8（容許 BOM）嚴格解碼；壞位元組 → format（不讓 U+FFFD 混進名稱）。"""
    if isinstance(value, (bytes, bytearray)):
        try:
            value = bytes(value).decode("utf-8-sig")
        except UnicodeDecodeError:
            raise _err("format", "encoding") from None
    if not isinstance(value, str):
        raise _err("format", "text")
    return value.lstrip("\ufeff")


def _as_json(payload: object) -> Any:
    """已解析的 JSON 原樣回傳；bytes／str 先解析。NaN／Infinity 當缺值（json 預設會收）。"""
    if isinstance(payload, (bytes, bytearray, str)):
        text = _as_text(payload)
        try:
            return json.loads(text, parse_constant=lambda _name: None)
        except (ValueError, RecursionError):
            raise _err("format", "json") from None
    return payload


def _get(obj: object, *keys: str) -> Any:
    """一路往下取 dict 的鍵；中間不是 dict 就回 None。"""
    for key in keys:
        if not isinstance(obj, dict):
            return None
        obj = obj.get(key)
    return obj


def _nasdaq_data(payload: object) -> dict:
    """Nasdaq API 外殼 {"data": {...}, "status": {...}}；data 是 null（例：代號不存在）→ empty。"""
    root = _as_json(payload)
    if not isinstance(root, dict) or "data" not in root:
        raise _err("format", "envelope")
    data = root["data"]
    if data is None:
        raise _err("empty", "data")
    if not isinstance(data, dict):
        raise _err("format", "data")
    return data


def _rows_list(value: object, field_name: str) -> list:
    """表格的 rows：null／[] → empty、不是 list → format。"""
    if value is None or value == []:
        raise _err("empty", field_name)
    if not isinstance(value, list):
        raise _err("format", field_name)
    return value


# ── 資料型別 ──────────────────────────────────────────────────────────────
@dataclass(frozen=True)
class Bar:
    """一根日線：Nasdaq 已調整分割的收盤與量（量是股數）；不還原股息。"""

    date: date
    close: float
    volume: Optional[int]
    open: Optional[float] = None
    high: Optional[float] = None
    low: Optional[float] = None


@dataclass(frozen=True)
class MarketInfo:
    """Nasdaq market-info：Nasdaq 認定的「今天」的收盤時刻（美東 aware）、前／下一個交易日。"""

    closing_time: datetime
    previous_trade_date: date
    next_trade_date: date
    is_business_day: bool
    opening_time: Optional[datetime] = None


@dataclass(frozen=True)
class Split:
    """分割事件：ratio「a : b」＝a 股新股換 b 股舊股；「2%」股票股利記成 1.02 : 1。

    比例看不懂時 new／old 是 None，但事件照樣回傳——它只用來觸發整段重抓（PLAN §3.1、§18.6）。"""

    symbol: str
    execution_date: date
    new: Optional[float]
    old: Optional[float]
    ratio_text: str = ""

    @property
    def factor(self) -> Optional[float]:
        """新股數／舊股數（3 : 1 → 3.0、1:10 → 0.1）；比例未知回 None。"""
        if self.new is None or self.old is None:
            return None
        return self.new / self.old


@dataclass(frozen=True)
class Holding:
    """持股一列。weight 是百分比（7.88 代表 7.88%）；Nasdaq-100 名單沒有權重 → None、改給市值（美元）。"""

    symbol: str
    name: str
    weight: Optional[float]
    market_cap: Optional[float] = None


@dataclass(frozen=True)
class HoldingList:
    """持股清單。rows 只有可用的股票列（canonical 代號），依來源順序；同代號重複的權重相加。

    dropped_count：非股票列與壞列（現金、期貨、擔保品、CVR、代號或權重不合格）。
    unmapped_*：是股票但對不到美國代號（例 SK Hynix）——卡上寫「另有 N 檔無美國代號（合計 X%）」；
    unmapped_names 已過 market_calc.clean_external（≤24 字，最多 20 個）。rows 的 name 只做正規化，
    進 SQLite 前由 job 清洗。
    可以直接 `for h in holding_list`、`len(holding_list)`。"""

    as_of: Optional[date]
    rows: list[Holding]
    dropped_count: int
    unmapped_weight: float
    unmapped_count: int = 0
    unmapped_names: tuple[str, ...] = ()

    def __iter__(self) -> Iterator[Holding]:
        return iter(self.rows)

    def __len__(self) -> int:
        return len(self.rows)

    @property
    def covered_weight(self) -> float:
        """有代號的股票列權重合計（SOXX 名單依實際涵蓋權重標示）。"""
        return round(sum(h.weight for h in self.rows if h.weight is not None), 4)


@dataclass(frozen=True)
class OptionContract:
    """一個期權合約。kind 是 "C"／"P"（和代號同一個字母）；strike 是美元；iv 是小數（0.2267）。

    這裡只做「格式有效」：代號合格、OI 是非負整數、履約價 >0、到期日合法。剩餘期限 >0、gamma 有效、
    IV=0 改用價外 IV 都要看資料日，由 market_calc.option_sets 決定。"""

    symbol: str
    expiry: date
    kind: str
    strike: float
    oi: int
    volume: Optional[int]
    iv: Optional[float]
    gamma: Optional[float]
    delta: Optional[float]


@dataclass(frozen=True)
class OptionConflict:
    """同一個代號出現多筆且值不同（兩筆都不用）；oi 只有各筆一致時才知道，否則 None。"""

    symbol: str
    expiry: date
    kind: str
    strike: float
    oi: Optional[int]
    rows: int


@dataclass(frozen=True)
class OptionChain:
    """CBOE 期權鏈。timestamp_utc 是檔案產生時間（只用來拒收：抓取時間−它 >12 小時）；資料日是
    quote_last_trade_et 的美東日期（PLAN §18.1）。spot／close 是 CBOE 的最新成交與收盤（美元）、iv30 是
    百分比、volume 是標的當日量。raw_count＝收到的列數；bad_count＝格式無效；conflict_count＝因同代號
    值衝突排除的列數（詳細在 conflicts）；conflict_oi_known＝每個衝突代號的 OI 都知道（覆蓋率不是下限）。"""

    timestamp_utc: datetime
    spot: float
    iv30: Optional[float]
    contracts: list[OptionContract]
    raw_count: int
    bad_count: int
    conflict_count: int
    conflict_oi_known: bool
    quote_last_trade_et: datetime
    symbol: str
    close: Optional[float] = None
    volume: Optional[int] = None
    conflicts: tuple[OptionConflict, ...] = ()


@dataclass(frozen=True)
class Quote:
    """CBOE 指數報價（VIX、VIX3M）：日變化用 price_change（收盤後 prev_day_close 會等於 close）。"""

    close: float
    price_change: Optional[float]
    price_change_pct: Optional[float]
    last_trade_et: datetime
    timestamp_utc: datetime
    symbol: str = ""


@dataclass(frozen=True)
class YieldRow:
    """美國財政部日殖利率（%）：3 個月、2 年、10 年、30 年；某天沒公布的天期是 None。"""

    date: date
    m3: Optional[float]
    y2: Optional[float]
    y10: Optional[float]
    y30: Optional[float]


@dataclass(frozen=True)
class CftcRow:
    """CFTC TFF（期貨）一列；口數。chg_* 是報告裡的週變化欄（沒有就是 None）。"""

    market: str
    report_date: date
    oi: int
    dealer_l: int
    dealer_s: int
    am_l: int
    am_s: int
    lev_l: int
    lev_s: int
    other_l: int
    other_s: int
    code: str = ""
    nonrept_l: Optional[int] = None
    nonrept_s: Optional[int] = None
    chg_oi: Optional[int] = None
    chg_dealer_l: Optional[int] = None
    chg_dealer_s: Optional[int] = None
    chg_am_l: Optional[int] = None
    chg_am_s: Optional[int] = None
    chg_lev_l: Optional[int] = None
    chg_lev_s: Optional[int] = None
    chg_other_l: Optional[int] = None
    chg_other_s: Optional[int] = None

    @property
    def lev_net(self) -> int:
        return self.lev_l - self.lev_s

    @property
    def am_net(self) -> int:
        return self.am_l - self.am_s

    @property
    def lev_net_change(self) -> Optional[int]:
        if self.chg_lev_l is None or self.chg_lev_s is None:
            return None
        return self.chg_lev_l - self.chg_lev_s

    @property
    def am_net_change(self) -> Optional[int]:
        if self.chg_am_l is None or self.chg_am_s is None:
            return None
        return self.chg_am_l - self.chg_am_s


@dataclass(frozen=True)
class CftcPair:
    """S&P 500 與 Nasdaq-100 兩個合約，一定是同一種口徑（basis：consolidated 或 e-mini）。"""

    es: CftcRow
    nq: CftcRow
    basis: str


@dataclass(frozen=True)
class MarginRow:
    """FINRA 融資餘額：month 是該月 1 日；debit_musd 是百萬美元。"""

    month: date
    debit_musd: int


@dataclass(frozen=True)
class Inst13F:
    """Nasdaq 13F 彙總：own_pct 是機構持股率（%，沒有就 None）；其他是機構家數。holder_periods 是前十大
    持有人各自的申報期間（依來源順序；job 取眾數，PLAN §18.9）。"""

    own_pct: Optional[float]
    increased: int
    decreased: int
    held: int
    new: int
    soldout: int
    holder_periods: list[date]
    total_holders: Optional[int] = None


@dataclass(frozen=True)
class Quarter:
    """季度損益：期末日與「Net Income Applicable to Common Shareholders」（千美元，沒有就 None）。"""

    end_date: date
    net_income_common_kusd: Optional[float]


@dataclass(frozen=True)
class FomcMeeting:
    """FOMC 會期（美東日期）；sep＝有經濟預測摘要（頁面上的「*」）。"""

    start: date
    end: date
    sep: bool = False

    @property
    def decision_utc(self) -> datetime:
        """利率決議時刻：會期最後一天 14:00 美東（PLAN §18.9），換成 UTC。"""
        return datetime.combine(self.end, time(14, 0), tzinfo=_ny()).astimezone(UTC)


@dataclass(frozen=True)
class EconEvent:
    """經濟事件：name 是固定白名單裡的英文名（不是來源原字串）；中文標籤查 US_EVENT_LABELS。"""

    time_utc: datetime
    country: str
    name: str


@dataclass(frozen=True)
class EarningsDate:
    """財報日；estimated＝來源寫明是演算法估的（卡上一律標「預估」）。"""

    date: date
    estimated: bool


# ── Nasdaq：日線、market-info、分割、Nasdaq-100、13F、基本面、事件、財報日 ────────
def parse_nasdaq_historical(payload: object, *, expect_symbol: Optional[str] = None) -> list[Bar]:
    """/api/quote/{SYM}/historical：個股收盤帶「$」、ETF 不帶，量有千分位；日期 MM/DD/YYYY。由舊到新。"""
    data = _nasdaq_data(payload)
    if expect_symbol is not None and canonical_us_symbol(data.get("symbol"), "nasdaq") != expect_symbol:
        raise _err("format", "symbol")
    table = data.get("tradesTable")
    if table is None:
        raise _err("empty", "tradesTable")
    if not isinstance(table, dict):
        raise _err("format", "tradesTable")
    bars: dict[date, Bar] = {}
    for row in _rows_list(table.get("rows"), "rows"):
        if not isinstance(row, dict):
            raise _err("format", "row")
        day = _mdy(row.get("date"), "date")
        close = _num(row.get("close"))
        if close is None:
            raise _err("missing_field", "close")
        if close <= 0:
            raise _err("format", "close")
        volume = _int(row.get("volume"))
        if volume is not None and volume < 0:
            raise _err("format", "volume")
        bar = Bar(
            day,
            close,
            volume,
            _positive_or_none(row.get("open")),
            _positive_or_none(row.get("high")),
            _positive_or_none(row.get("low")),
        )
        previous = bars.get(day)
        if previous is not None and previous != bar:
            raise _err("conflict", "date")
        bars[day] = bar
    return [bars[day] for day in sorted(bars)]


_MARKET_TIME_RE = re.compile(
    r"([A-Za-z]{3})[A-Za-z]*\.?\s+(\d{1,2}),\s*(\d{4})\s+(\d{1,2}):(\d{2})\s*([AaPp][Mm])(?:\s*ET)?"
)


def _market_time(text: object, raw: object, field_name: str) -> Optional[datetime]:
    """「Oct 16, 2026 04:00 PM ET」與「2026-10-16T16:00:00」（都是美東當地時間）；兩個都有就要一致。"""
    parsed: list[datetime] = []
    if not _is_blank(text):
        match = _MARKET_TIME_RE.fullmatch(text.strip()) if isinstance(text, str) else None
        month = _MONTHS.get(match.group(1).lower()) if match else None
        if not match or month is None:
            raise _err("date", field_name)
        hour, minute = int(match.group(4)), int(match.group(5))
        if not (1 <= hour <= 12 and minute < 60):
            raise _err("date", field_name)
        hour = hour % 12 + (12 if match.group(6).lower() == "pm" else 0)
        day = _safe_date(int(match.group(3)), month, int(match.group(2)), field_name)
        parsed.append(datetime.combine(day, time(hour, minute), tzinfo=_ny()))
    if not _is_blank(raw):
        parsed.append(_aware(raw, _ny(), field_name))
    if not parsed:
        return None
    if any(value != parsed[0] for value in parsed[1:]):
        raise _err("conflict", field_name)
    return parsed[0]


def parse_nasdaq_market_info(payload: object) -> MarketInfo:
    """/api/market-info：給每日維護核對 NYSE 表（PLAN §18.3；以美東今天為準的核對由 job 做）。"""
    data = _nasdaq_data(payload)
    closing = _market_time(data.get("marketClosingTime"), data.get("closeRaw"), "marketClosingTime")
    if closing is None:
        raise _err("missing_field", "marketClosingTime")
    try:
        opening = _market_time(data.get("marketOpeningTime"), data.get("openRaw"), "marketOpeningTime")
    except ParseError:
        opening = None  # 開盤時刻只是附帶資訊，壞了不擋核對
    previous = _mon_d_y(data.get("previousTradeDate"), "previousTradeDate")
    following = _mon_d_y(data.get("nextTradeDate"), "nextTradeDate")
    flag = data.get("isBusinessDay")
    if flag is None:
        raise _err("missing_field", "isBusinessDay")
    if not isinstance(flag, bool):
        raise _err("format", "isBusinessDay")
    if not previous < following:
        raise _err("conflict", "tradeDates")
    return MarketInfo(closing, previous, following, flag, opening)


_RATIO_RE = re.compile(r"(\d+(?:\.\d+)?)\s*:\s*(\d+(?:\.\d+)?)")
_PERCENT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%")
_RATIO_TEXT_OK = re.compile(r"[0-9 .:%]{1,32}")


def _split_ratio(value: object) -> tuple[Optional[float], Optional[float], str]:
    """「3 : 1」「1:10」「1 : 12.62」「1.5:1」→ (新, 舊)；「2%」股票股利 → (1.02, 1.0)；看不懂 → (None, None)。"""
    text = _norm_text(value, 32)
    shown = text if _RATIO_TEXT_OK.fullmatch(text) else ""
    match = _RATIO_RE.fullmatch(text)
    if match:
        new, old = float(match.group(1)), float(match.group(2))
        if new > 0 and old > 0:
            return new, old, shown
        return None, None, shown
    match = _PERCENT_RE.fullmatch(text)
    if match and float(match.group(1)) > 0:
        return 1.0 + float(match.group(1)) / 100.0, 1.0, shown
    return None, None, shown


def parse_nasdaq_splits(payload: object) -> list[Split]:
    """/api/calendar/splits?date=X：executionDate ≥X 的全部事件（實測 X＝6/1 → 364 筆）。

    沒有事件是正常的（回 []）。代號不是 canonical 的列（權證、外國代號）不可能在成分裡，略過；日期壞掉的
    列略過，但整份都壞 → date。同一檔同一天兩個不同比例 → 比例當未知（事件照樣觸發重抓）。"""
    data = _nasdaq_data(payload)
    rows = data.get("rows")
    if rows is None:
        return []
    if not isinstance(rows, list):
        raise _err("format", "rows")
    found: dict[tuple[str, date], Split] = {}
    bad_dates = 0
    for row in rows:
        if not isinstance(row, dict):
            raise _err("format", "row")
        symbol = canonical_us_symbol(row.get("symbol"), "nasdaq")
        if symbol is None:
            continue
        try:
            day = _mdy(row.get("executionDate"), "executionDate")
        except ParseError:
            bad_dates += 1
            continue
        new, old, text = _split_ratio(row.get("ratio"))
        split = Split(symbol, day, new, old, text)
        previous = found.get((symbol, day))
        if previous is not None and previous != split:
            split = Split(symbol, day, None, None, "")
        found[(symbol, day)] = split
    if bad_dates and not found:
        raise _err("date", "executionDate")
    return sorted(found.values(), key=lambda s: (s.execution_date, s.symbol))


def parse_nasdaq100(payload: object, *, min_count: int = 95) -> HoldingList:
    """/api/quote/list-type/nasdaq100：成分（約 101 檔，GOOGL／GOOG 各一）＋市值；沒有權重。

    回 HoldingList（INTERFACES 寫 list[Holding]；HoldingList 可直接迭代、並多帶名單日期）。"""
    data = _nasdaq_data(payload)
    rows = _rows_list(_get(data, "data", "rows"), "rows")
    holdings: dict[str, Holding] = {}
    dropped = 0
    for row in rows:
        if not isinstance(row, dict):
            raise _err("format", "row")
        symbol = canonical_us_symbol(row.get("symbol"), "nasdaq")
        if symbol is None:
            dropped += 1
            continue
        cap = _positive_or_none(row.get("marketCap"))
        holding = Holding(symbol, _norm_text(row.get("companyName")), None, cap)
        if symbol in holdings and holdings[symbol] != holding:
            raise _err("conflict", "symbol")
        holdings[symbol] = holding
    total = data.get("totalrecords")
    if total is not None and _num_or_none(total) != float(len(rows)):
        raise _err("too_few", "totalrecords")
    if len(holdings) < min_count:
        raise _err("too_few", "rows")
    stamp = data.get("date")
    match = _MON_D_Y_RE.search(stamp) if isinstance(stamp, str) else None
    as_of = None
    if match and _MONTHS.get(match.group(1).lower()):
        as_of = _safe_date(int(match.group(3)), _MONTHS[match.group(1).lower()], int(match.group(2)), "date")
    return HoldingList(as_of, list(holdings.values()), dropped, 0.0)


def _label_counts(rows: object) -> dict[str, Optional[int]]:
    """13F 的「Increased Positions／3,369」這種列 → {"increased positions": 3369}。"""
    out: dict[str, Optional[int]] = {}
    if not isinstance(rows, list):
        return out
    for row in rows:
        if not isinstance(row, dict):
            continue
        label = _norm_text(row.get("positions")).lower()
        if label:
            out[label] = _count_or_none(row.get("holders"))
    return out


def parse_nasdaq_13f(payload: object) -> Inst13F:
    """/api/company/{SYM}/institutional-holdings：加碼／減碼／持平／新建／出清的機構家數、機構持股率、
    前十大持有人的申報期間（實測 Vanguard 12/31/2025、其他多是 6/30/2026，所以是混合期間）。"""
    data = _nasdaq_data(payload)
    active = _label_counts(_get(data, "activePositions", "rows"))
    newsold = _label_counts(_get(data, "newSoldOutPositions", "rows"))
    wanted = {
        "increased": active.get("increased positions"),
        "decreased": active.get("decreased positions"),
        "held": active.get("held positions"),
        "new": newsold.get("new positions"),
        "soldout": newsold.get("sold out positions"),
    }
    if all(value is None for value in wanted.values()):
        raise _err("empty", "positions")
    missing = [name for name, value in wanted.items() if value is None]
    if missing:
        raise _err("missing_field", missing[0])
    own_pct = _num_or_none(_get(data, "ownershipSummary", "SharesOutstandingPCT", "value"))
    if own_pct is not None and not 0 <= own_pct <= 1000:
        own_pct = None  # 持股率可以 >100%（重複申報），但離譜的值不顯示
    periods: list[date] = []
    holders = _get(data, "holdingsTransactions", "table", "rows")
    if isinstance(holders, list):
        for row in holders:
            if isinstance(row, dict):
                try:
                    periods.append(_mdy(row.get("date"), "date"))
                except ParseError:
                    continue
    return Inst13F(
        own_pct,
        wanted["increased"],
        wanted["decreased"],
        wanted["held"],
        wanted["new"],
        wanted["soldout"],
        periods,
        active.get("total institutional shares"),
    )


_NET_INCOME_COMMON = "net income applicable to common shareholders"
_PERIOD_LABEL_RE = re.compile(r"(quarter|period)\s+ending", re.IGNORECASE)


def parse_nasdaq_financials(payload: object) -> list[Quarter]:
    """/api/company/{SYM}/financials?frequency=2 的季度損益表（PLAN §19.3）：近四季普通股淨利（千美元）。

    沒有存下真實回應當 fixture；依 Nasdaq 這個端點的已知形狀寫（incomeStatementTable 的 headers／rows
    用 value1…valueN 當欄，期末日在表頭或「Quarter Ending:」那一列）。對不上 → format／missing_field，
    job 讓那格暫缺（不走 Yahoo）。"""
    data = _nasdaq_data(payload)
    table = data.get("incomeStatementTable")
    if table is None:
        raise _err("empty", "incomeStatementTable")
    if not isinstance(table, dict):
        raise _err("format", "incomeStatementTable")
    rows = _rows_list(table.get("rows"), "rows")
    if not all(isinstance(row, dict) for row in rows):
        raise _err("format", "row")
    headers = table.get("headers") if isinstance(table.get("headers"), dict) else {}
    columns = sorted(
        (key for key in set(headers) | {k for row in rows for k in row} if re.fullmatch(r"value([2-9]|\d{2})", key)),
        key=lambda key: int(key[5:]),
    )
    ends: dict[str, date] = {}
    for source in [headers] + [row for row in rows if _PERIOD_LABEL_RE.search(_norm_text(row.get("value1")))]:
        candidate: dict[str, date] = {}
        for key in columns:
            if _is_blank(source.get(key)):
                continue
            try:
                candidate[key] = _mdy(source.get(key), "period")
            except ParseError:
                candidate = {}
                break
        if candidate:
            ends = candidate
            break
    if not ends:
        raise _err("format", "period")
    if len(set(ends.values())) != len(ends):
        raise _err("conflict", "period")
    target = next((row for row in rows if _norm_text(row.get("value1")).lower() == _NET_INCOME_COMMON), None)
    if target is None:
        raise _err("missing_field", "net_income_common")
    quarters = [Quarter(end, _num(target.get(key))) for key, end in ends.items()]
    return sorted(quarters, key=lambda q: q.end_date)


_MARKET_CAP_RE = re.compile(r"\$?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?")
_MARKET_CAP_MAX = 1e14  # 100 兆美元：再大就是欄位壞了


def parse_nasdaq_summary_market_cap(payload: object, *, expect_symbol: Optional[str] = None) -> Optional[float]:
    """/api/quote/{SYM}/summary?assetclass=stocks 的市值（美元；PLAN §19.3「nasdaq100 名單或 summary」）。

    `summaryData.MarketCap.value` 是千分位整數字串（JPM「903,039,014,165」），是整家公司的市值（和 Nasdaq-100
    名單一樣：GOOG／GOOGL 各回整家公司的值）。「N/A」、沒有這欄、0 → None（不是 0）；負數、百分比、看不懂
    的字、大到不合理 → format。回應的代號用「/」分股別（請求 BRK.B、回 BRK/B），canonical 之後再和請求的比。"""
    data = _nasdaq_data(payload)
    if expect_symbol is not None and canonical_us_symbol(data.get("symbol"), "nasdaq") != expect_symbol:
        raise _err("format", "symbol")
    summary = data.get("summaryData")
    if summary is None:
        raise _err("empty", "summaryData")
    if not isinstance(summary, dict):
        raise _err("format", "summaryData")
    entry = summary.get("MarketCap")
    if entry is None:
        return None
    if not isinstance(entry, dict):
        raise _err("format", "MarketCap")
    value = entry.get("value")
    if _is_blank(value) or (isinstance(value, str) and value.strip().upper() in _MISSING_TOKENS):
        return None
    if not isinstance(value, str) or not _MARKET_CAP_RE.fullmatch(value.strip()):
        raise _err("format", "MarketCap")
    cap = _num(value)
    if cap is None or cap <= 0:
        return None
    if cap > _MARKET_CAP_MAX:
        raise _err("format", "MarketCap")
    return cap


# 經濟日曆白名單（Nasdaq 轉載 investing.com 的事件名；比對前去掉尾巴的「(MoM)」之類）→ 固定英文名。
# FOMC 會議本身用 Fed 頁（parse_fed_fomc），這裡不重複列利率決議。
US_EVENT_NAMES: dict[str, str] = {
    "cpi": "CPI",
    "core cpi": "Core CPI",
    "ppi": "PPI",
    "core ppi": "Core PPI",
    "nonfarm payrolls": "Nonfarm Payrolls",
    "unemployment rate": "Unemployment Rate",
    "gdp": "GDP",
    "pce price index": "PCE Price Index",
    "core pce price index": "Core PCE Price Index",
    "retail sales": "Retail Sales",
    "core retail sales": "Core Retail Sales",
    "ism manufacturing pmi": "ISM Manufacturing PMI",
    "ism non-manufacturing pmi": "ISM Services PMI",
    "ism services pmi": "ISM Services PMI",
    "initial jobless claims": "Initial Jobless Claims",
    "jolts job openings": "JOLTS Job Openings",
    "michigan consumer sentiment": "Michigan Consumer Sentiment",
    "fomc meeting minutes": "FOMC Minutes",
}
_FED_CHAIR_RE = re.compile(r"fed chair \w+(?: \w+)? speaks")
# 卡片用的標籤（PLAN §6）：表放在 market_calc，卡片程式也直接查同一份（2026-10-18）。
US_EVENT_LABELS = market_calc.US_EVENT_LABELS


_ECON_WEDNESDAY_SENTINEL = "mba mortgage applications"  # 每週三 07:00（美東）


def _us_event_name(value: object) -> Optional[str]:
    text = _norm_text(value, 120).lower()
    text = re.sub(r"\s*\([^()]*\)\s*$", "", text).strip()  # 「CPI (MoM)」→「cpi」
    if text in US_EVENT_NAMES:
        return US_EVENT_NAMES[text]
    if _FED_CHAIR_RE.fullmatch(text):
        return "Fed Chair Speaks"
    return None


def parse_nasdaq_econ_calendar(payload: object, day: date, *, require_sentinel: bool = False) -> list[EconEvent]:
    """美東 `day` 這天的經濟事件（請求要用 date=day＋1，見 fetch_nasdaq_econ_calendar）：只留美國白名單，
    `gmt` 欄（美東當地時刻）＋day → UTC。

    同一時刻同名（例：Core CPI 年增與月增各一列）合併成一則；沒有時刻（「All Day」）的不列。
    日期對不對用每週三固定公布的「MBA Mortgage Applications」驗：它出現在別的星期幾，就代表這份回應的
    日期對應錯了（例如 Nasdaq 修好了差一天的問題）→ conflict，這一天不收，任務記失敗、每日維護看得到。"""
    data = _nasdaq_data(payload)
    rows = data.get("rows")
    if rows is None:
        return []
    if not isinstance(rows, list):
        raise _err("format", "rows")
    events: dict[tuple[datetime, str], EconEvent] = {}
    sentinel_seen = False
    for row in rows:
        if not isinstance(row, dict):
            raise _err("format", "row")
        if _norm_text(row.get("country")) != "United States":
            continue
        if _norm_text(row.get("eventName"), 120).lower() == _ECON_WEDNESDAY_SENTINEL:
            sentinel_seen = True
            if day.weekday() != 2:
                raise _err("conflict", "weekday")
        name = _us_event_name(row.get("eventName"))
        clock = row.get("gmt")
        match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})\s*", clock) if isinstance(clock, str) else None
        if name is None or match is None or int(match.group(1)) > 23 or int(match.group(2)) > 59:
            continue
        local = datetime.combine(day, time(int(match.group(1)), int(match.group(2))), tzinfo=_ny())
        when = local.astimezone(UTC)
        events[(when, name)] = EconEvent(when, "US", name)
    # 2026-10-18 審查 r6b I1：只有出現在別天才擋的話，Nasdaq 修好位移時只有一天會錯、其他六天整批早一天寫進去；
    # 所以呼叫端用週三驗一次：週三一定要看到 MBA（require_sentinel），沒看到就整批不收
    if require_sentinel and not sentinel_seen:
        raise _err("conflict", "weekday")
    return [events[key] for key in sorted(events)]


_EARNINGS_TEXT_DATE = re.compile(r"\b(\d{1,2}/\d{1,2}/\d{4})\b")


def parse_nasdaq_earnings_date(payload: object, *, expect_symbol: Optional[str] = None) -> EarningsDate:
    """/api/analyst/{SYM}/earnings-date：reportText 的 MM/DD/YYYY 與 announcement 的「Nov 18, 2026」要一致。"""
    data = _nasdaq_data(payload)
    text = data.get("reportText") if isinstance(data.get("reportText"), str) else ""
    announcement = data.get("announcement") if isinstance(data.get("announcement"), str) else ""
    found: list[date] = []
    match = _EARNINGS_TEXT_DATE.search(text)
    if match:
        found.append(_mdy(match.group(1), "reportText"))
    tail = announcement.rsplit(":", 1)[-1] if ":" in announcement else ""
    match = _MON_D_Y_RE.search(tail)
    if match:
        found.append(_mon_d_y(match.group(0), "announcement"))
    if expect_symbol is not None:
        named = re.search(r"\bfor\s+([A-Z][A-Z./]{0,7})\s*:", announcement)
        if named and canonical_us_symbol(named.group(1), "nasdaq") != expect_symbol:
            raise _err("format", "symbol")
    if not found:
        raise _err("empty", "earnings_date")
    if any(value != found[0] for value in found[1:]):
        raise _err("conflict", "earnings_date")
    lowered = text.lower()
    estimated = "estimated" in lowered or "algorithm" in lowered or "*" in announcement
    return EarningsDate(found[0], estimated)


# ── 持股表共用 ────────────────────────────────────────────────────────────
_NON_EQUITY_RE = re.compile(r"\b(CASH|CSH|FUTURES?|SWAPS?|COLLATERAL|MONEY MARKET)\b")
_NON_EQUITY_ZH = ("期貨", "現金", "保證金", "選擇權")


def _is_non_equity_name(name: str) -> bool:
    """現金、期貨、交換、擔保品、貨幣市場列（PLAN §3.1）；用整個字比，「Dollar General」不會被誤殺。"""
    return bool(_NON_EQUITY_RE.search(_match_key(name))) or any(word in name for word in _NON_EQUITY_ZH)


def _cusip_ok(value: object) -> bool:
    """9 碼 CUSIP（含檢查碼）。SSGA 的現金列也有合格的 CUSIP，所以另外還要看代號。"""
    if not isinstance(value, str):
        return False
    text = value.strip().upper()
    if not _CUSIP_RE.fullmatch(text):
        return False
    total = 0
    for index, ch in enumerate(text[:8]):
        if ch.isdigit():
            digit = int(ch)
        elif "A" <= ch <= "Z":
            digit = ord(ch) - 55
        else:
            digit = {"*": 36, "@": 37, "#": 38}[ch]
        if index % 2:
            digit *= 2
        total += digit // 10 + digit % 10
    return (10 - total % 10) % 10 == int(text[8])


def _stored_name(text: str, max_len: int = 24) -> Optional[str]:
    """要原樣存進 SQLite 的外部名稱先過共用清洗 `market_calc.clean_external`（PLAN §19.1）；
    清洗後是 None（含網址、email、手機號）就丟掉；那個模組不在時一律不存。"""
    try:
        from market_calc import clean_external  # noqa: PLC0415
    except ImportError:
        return None
    return clean_external(text, max_len)


class _HoldingsBuilder:
    """收持股列：同代號權重相加、記非股票列與對不到代號的權重。"""

    def __init__(self) -> None:
        self.rows: dict[str, Holding] = {}
        self.dropped = 0
        self.unmapped_weight = 0.0
        self.unmapped_count = 0
        self.unmapped_names: list[str] = []

    def add(self, symbol: str, name: str, weight: float) -> None:
        old = self.rows.get(symbol)
        if old is not None:
            weight += old.weight or 0.0
            name = old.name
        self.rows[symbol] = Holding(symbol, name, round(weight, 6))

    def unmapped(self, name: str, weight: float) -> None:
        self.unmapped_count += 1
        self.unmapped_weight += weight
        # 名稱只給每日維護補名稱→代號表用；job 會把整份清單（含這欄）存進 SQLite，所以這裡就先清洗
        cleaned = _stored_name(name)
        if cleaned is not None and len(self.unmapped_names) < 20:
            self.unmapped_names.append(cleaned)

    def result(self, as_of: Optional[date]) -> HoldingList:
        return HoldingList(
            as_of,
            list(self.rows.values()),
            self.dropped,
            round(self.unmapped_weight, 4),
            self.unmapped_count,
            tuple(self.unmapped_names),
        )


# ── xlsx（SSGA、FINRA）────────────────────────────────────────────────────
XLSX_MAX_MEMBERS = 64
XLSX_MAX_UNCOMPRESSED = 50 * MIB
XLSX_MAX_ROWS = 20000
_XML_DECL_RE = re.compile(rb"<!\s*(doctype|entity)", re.IGNORECASE)
_XML_MEMBER_SUFFIXES = (".xml", ".rels", ".vml")
_WIDE_BOMS = (b"\xff\xfe", b"\xfe\xff", b"\x00\x00\xfe\xff")


def _xlsx_guard(data: object) -> bytes:
    """交給 openpyxl 之前先看 zip 目錄與每個 XML 成員（PLAN §19.2）：

    成員 ≤64、宣告的解壓總量 ≤50 MB（zipfile 讀成員時不會超過宣告大小，CRC 也會驗）；XML 成員不得有
    DOCTYPE／ENTITY（實體展開與外部實體），也不得是 UTF-16／UTF-32——那樣的位元組掃描會看不到宣告。"""
    if not isinstance(data, (bytes, bytearray)):
        raise _err("format", "xlsx")
    if not data:
        raise _err("empty", "xlsx")
    try:
        archive = zipfile.ZipFile(io.BytesIO(bytes(data)))
    except (zipfile.BadZipFile, ValueError, OSError):
        raise _err("format", "zip") from None
    with archive:
        members = archive.infolist()
        if not members or len(members) > XLSX_MAX_MEMBERS:
            raise _err("format", "zip_members")
        if sum(info.file_size for info in members) > XLSX_MAX_UNCOMPRESSED:
            raise _err("format", "zip_size")
        for info in members:
            if info.is_dir():
                continue
            try:
                content = archive.read(info)
            except (zipfile.BadZipFile, ValueError, OSError, NotImplementedError, RuntimeError):
                raise _err("format", "zip_member") from None
            # 2026-10-18 安全審查 I4：不能只看副檔名或開頭是不是「<?xml」——openpyxl 照 rels 找成員，改名成 .dat、
            # 前面加 BOM／空白都照樣會被當 XML 讀。每個成員去掉 BOM 與空白後，以「<」開頭（或像沒有 BOM 的 UTF-16 XML）
            # 就一律檢查；真正的二進位成員（例 SSGA 的 xl/printerSettings/printerSettings1.bin，內含 NUL）才略過。
            head = content.lstrip(b"\xef\xbb\xbf").lstrip(b" \t\r\n")
            wide = content.startswith(_WIDE_BOMS) or head[:2] in (b"<\x00", b"\x00<")
            looks_xml = info.filename.lower().endswith(_XML_MEMBER_SUFFIXES) or head[:1] == b"<" or wide
            if not looks_xml:
                continue
            if wide or b"\x00" in content:
                raise _err("format", "xml_encoding")
            if _XML_DECL_RE.search(content):
                raise _err("format", "xml_doctype")
    return bytes(data)


def _xlsx_rows(data: object) -> list[tuple]:
    """第一張工作表的全部列（values_only）。先過 `_xlsx_guard`；openpyxl 的任何失敗都算 format。"""
    checked = _xlsx_guard(data)
    import openpyxl  # noqa: PLC0415 - 只有解析 xlsx 時才載入

    rows: list[tuple] = []
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = openpyxl.load_workbook(io.BytesIO(checked), read_only=True, data_only=True)
            try:
                if not workbook.worksheets:
                    raise _err("format", "worksheet")
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()  # 有些檔的 dimension 寫錯，read_only 會因此截掉欄
                for row in sheet.iter_rows(values_only=True):
                    rows.append(tuple(row))
                    if len(rows) > XLSX_MAX_ROWS:
                        raise _err("format", "rows")
            finally:
                workbook.close()
    except ParseError:
        raise
    except Exception:  # openpyxl 對壞檔會丟各式各樣的例外，全部當格式不對
        raise _err("format", "xlsx") from None
    return rows


def _cell(row: tuple, index: Optional[int]) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _cell_text(row: tuple, index: Optional[int]) -> str:
    value = _cell(row, index)
    if value is None:
        return ""
    return _norm_text(value if isinstance(value, str) else str(value), 200)


_SSGA_AS_OF_RE = re.compile(r"as of\s+(\d{1,2})-([A-Za-z]{3})-(\d{4})", re.IGNORECASE)


def _ssga_as_of(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = _SSGA_AS_OF_RE.search(value) if isinstance(value, str) else None
    month = _MONTHS.get(match.group(2).lower()) if match else None
    if not match or month is None:
        raise _err("date", "as_of")
    return _safe_date(int(match.group(3)), month, int(match.group(1)), "as_of")


def parse_ssga_spy_xlsx(data: bytes, *, min_rows: int = 450) -> HoldingList:
    """SSGA SPY 每日持股 xlsx（代表 VOO，Andrew 10/17 選）：「Holdings: As of 15-Oct-2026」＋表頭
    Name／Ticker／Identifier／Weight。非股票列（US DOLLAR、TPG 的 CVR「2602335D」、CUSIP 不合格、
    名稱有現金等字）濾掉記數；表格到第一個沒有代號、識別碼、權重的列為止（下面是免責聲明）。"""
    rows = _xlsx_rows(data)
    fund = None
    as_of: Optional[date] = None
    header_index = None
    for index, row in enumerate(rows[:40]):
        label = _cell_text(row, 0).lower()
        if label == "ticker symbol:":
            fund = _cell_text(row, 1).upper()
        elif label == "holdings:":
            as_of = _ssga_as_of(_cell(row, 1))
        elif label == "name" and "ticker" in [_cell_text(row, j).lower() for j in range(len(row))]:
            header_index = index
            break
    if header_index is None:
        raise _err("format", "header")
    if fund != "SPY":
        raise _err("format", "fund")
    if as_of is None:
        raise _err("date", "as_of")
    header = [_cell_text(rows[header_index], j).lower() for j in range(len(rows[header_index]))]
    columns = {name: header.index(name) for name in ("name", "ticker", "identifier", "weight") if name in header}
    if len(columns) != 4:
        raise _err("format", "columns")
    builder = _HoldingsBuilder()
    for row in rows[header_index + 1 :]:
        ticker, identifier = _cell(row, columns["ticker"]), _cell(row, columns["identifier"])
        raw_weight = _cell(row, columns["weight"])
        if _is_blank(ticker) and _is_blank(identifier) and _is_blank(raw_weight):
            break
        name = _cell_text(row, columns["name"])
        symbol = canonical_us_symbol(ticker, "ssga")
        weight = _num_or_none(raw_weight)
        if (
            symbol is None
            or not _cusip_ok(identifier)
            or _is_non_equity_name(name)
            or weight is None
            or weight < 0
        ):
            builder.dropped += 1
            continue
        builder.add(symbol, name, weight)
    if len(builder.rows) < min_rows:
        raise _err("too_few", "rows")
    return builder.result(as_of)


def _year_month(value: object) -> Optional[date]:
    """FINRA 的「2026-09」（或 Excel 日期）→ 該月 1 日；不像月份的列回 None（註腳、空白）。"""
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    match = re.fullmatch(r"\s*(\d{4})-(\d{1,2})\s*", value) if isinstance(value, str) else None
    if not match:
        return None
    return _safe_date(int(match.group(1)), int(match.group(2)), 1, "month")


def parse_finra_margin_xlsx(data: bytes, *, min_rows: int = 12) -> list[MarginRow]:
    """FINRA margin-statistics.xlsx：「Year-Month」＋「Debit Balances in Customers' Securities Margin
    Accounts」（百萬美元）。來源是新月份在上；回傳由舊到新。"""
    rows = _xlsx_rows(data)
    header_index = None
    for index, row in enumerate(rows[:20]):
        if _cell_text(row, 0).lower() == "year-month":
            header_index = index
            break
    if header_index is None:
        raise _err("format", "header")
    header = [_cell_text(rows[header_index], j).lower() for j in range(len(rows[header_index]))]
    debit_index = next((j for j, name in enumerate(header) if name.startswith("debit balances")), None)
    if debit_index is None:
        raise _err("format", "columns")
    months: dict[date, MarginRow] = {}
    for row in rows[header_index + 1 :]:
        month = _year_month(_cell(row, 0))
        if month is None:
            continue
        debit = _int(_cell(row, debit_index))
        if debit is None:
            raise _err("missing_field", "debit")
        if debit <= 0:
            raise _err("format", "debit")
        entry = MarginRow(month, debit)
        if month in months and months[month] != entry:
            raise _err("conflict", "month")
        months[month] = entry
    if len(months) < min_rows:
        raise _err("too_few", "rows")
    return [months[month] for month in sorted(months)]


# ── MoneyDJ 持股（HTML）───────────────────────────────────────────────────
class _Cell:
    __slots__ = ("texts", "hrefs", "header")

    def __init__(self, header: bool) -> None:
        self.texts: list[str] = []
        self.hrefs: list[str] = []
        self.header = header

    @property
    def text(self) -> str:
        return " ".join("".join(self.texts).split())


class _OpenTable:
    __slots__ = ("rows", "row", "cell")

    def __init__(self) -> None:
        self.rows: list[list[_Cell]] = []
        self.row: Optional[list[_Cell]] = None
        self.cell: Optional[_Cell] = None

    def end_cell(self) -> None:
        if self.cell is not None:
            if self.row is None:
                self.row = []
            self.row.append(self.cell)
            self.cell = None

    def end_row(self) -> None:
        self.end_cell()
        if self.row:
            self.rows.append(self.row)
        self.row = None


class _TableParser(HTMLParser):
    """把每個 <table> 收成「列 → 格（文字＋連結）」；巢狀表格各自一份。只讀文字與 href。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[_Cell]]] = []
        self._open: list[_OpenTable] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag == "table":
            table = _OpenTable()
            self.tables.append(table.rows)
            self._open.append(table)
            return
        if not self._open:
            return
        table = self._open[-1]
        if tag == "tr":
            table.end_row()
            table.row = []
        elif tag in ("td", "th"):
            table.end_cell()
            table.cell = _Cell(header=(tag == "th"))
        elif tag == "a" and table.cell is not None:
            href = dict(attrs).get("href")
            if href:
                table.cell.hrefs.append(href)
        elif tag == "br" and table.cell is not None:
            table.cell.texts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if not self._open:
            return
        table = self._open[-1]
        if tag in ("td", "th"):
            table.end_cell()
        elif tag == "tr":
            table.end_row()
        elif tag == "table":
            table.end_row()
            self._open.pop()

    def handle_data(self, data: str) -> None:
        if self._open and self._open[-1].cell is not None:
            self._open[-1].cell.texts.append(data)

    def finish(self) -> None:
        self.close()
        while self._open:
            self._open.pop().end_row()


def _html_tables(text: str) -> list[list[list[_Cell]]]:
    parser = _TableParser()
    try:
        parser.feed(text)
        parser.finish()
    except Exception:  # HTMLParser 遇到極端壞的標記可能丟例外：當格式不對
        raise _err("format", "html") from None
    return parser.tables


# MoneyDJ 沒有代號連結的列（多半是外國公司的 ADR 與股別股）：名稱開頭 → canonical 代號；None＝沒有美國
# 上市代號（PLAN §3.1、§16 實測 SOXX 的 8 檔外國股與 VOO 的 Berkshire、Brown-Forman）。
MONEYDJ_NAME_TO_SYMBOL: tuple[tuple[str, Optional[str]], ...] = (
    ("BERKSHIRE HATHAWAY", "BRK.B"),
    ("BROWN FORMAN", "BF.B"),
    ("TAIWAN SEMICONDUCTOR", "TSM"),
    ("ASML HOLDING", "ASML"),
    ("ARM HOLDINGS", "ARM"),
    ("UNITED MICRO ELECTRONICS", "UMC"),
    ("UNITED MICROELECTRONICS", "UMC"),
    ("ASE TECHNOLOGY", "ASX"),
    ("STMICROELECTRONICS", "STM"),
    ("TOWER SEMICONDUCTOR", "TSEM"),
    ("SK HYNIX", None),
)
_MONEYDJ_MIN_ROWS = {"SOXX": 20, "QQQ": 90, "VOO": 400, "SPY": 400, "0050": 40}
_MONEYDJ_DATE_RE = re.compile(r"資料日期\s*[:：]\s*(\d{4})/(\d{1,2})/(\d{1,2})")
_MONEYDJ_TITLE_RE = re.compile(r"\(([A-Z0-9.]{1,12})\)\s*-\s*全部持股")
# 只有帶市場尾碼的括號才是代號（「INTC.US」「2330.TW」「FITXN*1.TF」）；「(ADR)」這種是名稱的一部分。
_TRAILING_CODE_RE = re.compile(r"\(\s*([A-Za-z0-9*\-]{1,12}(?:\.[A-Za-z0-9]{1,4})*\.[A-Za-z]{2,3})\s*\)\s*$")


def _moneydj_etf(etf: object) -> tuple[str, str]:
    """「SOXX」→ ("SOXX", "us")；「0050」或「0050.TW」→ ("0050.TW", "tw")。其他寫法 → format。"""
    text = etf.strip().upper() if isinstance(etf, str) else ""
    if text.endswith(".TW"):
        text = text[:-3]
    if TW_CODE_RE.fullmatch(text):
        return text + ".TW", "tw"
    if US_SYMBOL_RE.fullmatch(text):
        return text, "us"
    raise _err("format", "etf")


def _moneydj_name_symbol(name: str) -> tuple[bool, Optional[str]]:
    """名稱表：(有沒有在表裡, 代號)。"""
    key = _match_key(name)
    for prefix, symbol in MONEYDJ_NAME_TO_SYMBOL:
        if key == prefix or key.startswith(prefix + " "):
            return True, symbol
    return False, None


def parse_moneydj_holdings(html_text: object, etf: str, *, min_rows: Optional[int] = None) -> HoldingList:
    """MoneyDJ「全部持股」頁（Basic0007B）：標題「…(SOXX)-全部持股」、「資料日期：2026/09/30」、
    表頭「個股名稱／投資比例(%)」。有連結的列用 etfid（美股「INTC.US」、台股「2330.TW」）；沒連結的查
    MONEYDJ_NAME_TO_SYMBOL；期貨（「.TF」、名稱有期貨）、現金、擔保品記成非股票列；對不到代號的股票
    （外國上市）記 unmapped。B（台股）也呼叫這個（etf="0050"）。"""
    etfid, market = _moneydj_etf(etf)
    text = _as_text(html_text)
    tables = _html_tables(text)
    table = None
    for candidate in tables:
        heads = [cell.text for row in candidate[:3] for cell in row if cell.header]
        if any("個股名稱" in h for h in heads) and any("投資比例" in h for h in heads):
            table = candidate
            break
    if table is None:
        raise _err("format", "table")
    backs = set()
    for row in table:
        for cell in row:
            for href in cell.hrefs:
                backs.update(value.upper() for value in parse_qs(urlsplit(href).query).get("back", []))
    titles = {match.group(1) for match in _MONEYDJ_TITLE_RE.finditer(text)}
    if etfid not in titles and not (backs and backs == {etfid}):
        raise _err("format", "etf")  # 不是要的那一檔（轉到別頁、首頁）
    match = _MONEYDJ_DATE_RE.search(text)
    if not match:
        raise _err("date", "as_of")
    as_of = _safe_date(int(match.group(1)), int(match.group(2)), int(match.group(3)), "as_of")
    builder = _HoldingsBuilder()
    for row in table:
        if len(row) < 2 or row[0].header:
            continue
        label = _norm_text(row[0].text, 120)
        weight = _num_or_none(row[1].text)
        code = None
        for href in row[0].hrefs:
            ids = parse_qs(urlsplit(href).query).get("etfid")
            if ids:
                code = ids[0].strip().upper()
                break
        trailing = _TRAILING_CODE_RE.search(label)
        if code is None and trailing:
            code = trailing.group(1).upper()
        name = _TRAILING_CODE_RE.sub("", label).strip() if trailing else label
        if not label or weight is None or weight < 0 or _is_non_equity_name(label):
            builder.dropped += 1
            continue
        if code is not None and code.endswith(".TF"):
            builder.dropped += 1  # 期貨
            continue
        symbol: Optional[str] = None
        if market == "us":
            if code is not None:
                symbol = canonical_us_symbol(code, "moneydj")
            else:
                _, symbol = _moneydj_name_symbol(name)
        elif code is not None and code.endswith(".TW") and TW_CODE_RE.fullmatch(code[:-3]):
            symbol = code[:-3]
        if symbol is None:
            builder.unmapped(name, weight)
        else:
            builder.add(symbol, name, weight)
    needed = min_rows if min_rows is not None else _MONEYDJ_MIN_ROWS.get(etfid.split(".")[0], 1)
    if len(builder.rows) + builder.unmapped_count < needed:
        raise _err("too_few", "rows")
    return builder.result(as_of)


# ── CBOE ─────────────────────────────────────────────────────────────────
def _cboe_contract(item: object, root: str) -> Optional[OptionContract]:
    """一筆合約；格式無效回 None（計入 bad_count）。"""
    if not isinstance(item, dict):
        return None
    code = item.get("option")
    match = OPTION_SYMBOL_RE.fullmatch(code.strip()) if isinstance(code, str) else None
    if not match or match.group(1) != root:
        return None
    ymd = match.group(2)
    try:
        expiry = date(2000 + int(ymd[:2]), int(ymd[2:4]), int(ymd[4:]))
    except ValueError:
        return None
    strike = int(match.group(4)) / 1000.0
    if strike <= 0:
        return None
    try:
        oi = _int(item.get("open_interest"))
    except ParseError:
        return None
    if oi is None or oi < 0:
        return None
    iv = _num_or_none(item.get("iv"))
    if iv is not None and iv < 0:
        iv = None
    return OptionContract(
        match.group(0),
        expiry,
        match.group(3),
        strike,
        oi,
        _count_or_none(item.get("volume")),
        iv,
        _num_or_none(item.get("gamma")),
        _num_or_none(item.get("delta")),
    )


def parse_cboe_options(payload: object, *, expect_symbol: Optional[str] = None) -> OptionChain:
    """cdn-api.cboe.com …/options/{SYM}.json。OI 只能是非負整數；IV=0 的合約照樣回傳（價內合約，
    market_calc 改用價外那邊的 IV）；同代號重複：值一樣合併、值不同兩筆都排除（記在 conflicts）。
    超過一半的列格式無效 → format（格式換了，不要硬算）。"""
    root = _as_json(payload)
    if not isinstance(root, dict):
        raise _err("format", "envelope")
    data = root.get("data")
    if data is None:
        raise _err("empty", "data")
    if not isinstance(data, dict):
        raise _err("format", "data")
    symbol = canonical_us_symbol(data.get("symbol") or root.get("symbol"), "cboe")
    if symbol not in CBOE_OPTION_SYMBOLS or (expect_symbol is not None and symbol != expect_symbol):
        raise _err("format", "symbol")
    timestamp_utc = _aware(root.get("timestamp"), UTC, "timestamp")
    last_trade_et = _aware(data.get("last_trade_time"), _ny(), "last_trade_time")
    spot = _positive_or_none(data.get("current_price")) or _positive_or_none(data.get("close"))
    if spot is None:
        raise _err("missing_field", "current_price")
    options = _rows_list(data.get("options"), "options")
    by_symbol: dict[str, list[OptionContract]] = {}
    bad = 0
    for item in options:
        contract = _cboe_contract(item, symbol)
        if contract is None:
            bad += 1
        else:
            by_symbol.setdefault(contract.symbol, []).append(contract)
    if bad * 2 > len(options):
        raise _err("format", "options")
    contracts: list[OptionContract] = []
    conflicts: list[OptionConflict] = []
    for group in by_symbol.values():
        first = group[0]
        if all(other == first for other in group[1:]):
            contracts.append(first)
            continue
        known = first.oi if all(other.oi == first.oi for other in group) else None
        conflicts.append(OptionConflict(first.symbol, first.expiry, first.kind, first.strike, known, len(group)))
    if not contracts:
        raise _err("too_few", "options")
    contracts.sort(key=lambda c: (c.expiry, c.strike, c.kind))
    conflicts.sort(key=lambda c: (c.expiry, c.strike, c.kind))
    iv30 = _num_or_none(data.get("iv30"))
    return OptionChain(
        timestamp_utc=timestamp_utc,
        spot=spot,
        iv30=iv30 if iv30 is not None and iv30 >= 0 else None,
        contracts=contracts,
        raw_count=len(options),
        bad_count=bad,
        conflict_count=sum(c.rows for c in conflicts),
        conflict_oi_known=all(c.oi is not None for c in conflicts),
        quote_last_trade_et=last_trade_et,
        symbol=symbol,
        close=_positive_or_none(data.get("close")),
        volume=_count_or_none(data.get("volume")),
        conflicts=tuple(conflicts),
    )


def _index_symbol(value: object) -> str:
    text = value.strip().upper().lstrip("^_") if isinstance(value, str) else ""
    return text if re.fullmatch(r"[A-Z0-9]{1,8}", text) else ""


def parse_cboe_quote(payload: object, *, expect_symbol: Optional[str] = None) -> Quote:
    """…/quotes/_VIX.json：資料日看 last_trade_time（美東）的日期；日變化用 price_change（PLAN §18.8）。"""
    root = _as_json(payload)
    if not isinstance(root, dict):
        raise _err("format", "envelope")
    data = root.get("data")
    if data is None:
        raise _err("empty", "data")
    if not isinstance(data, dict):
        raise _err("format", "data")
    symbol = _index_symbol(data.get("symbol")) or _index_symbol(root.get("symbol"))
    if not symbol or (expect_symbol is not None and symbol != expect_symbol):
        raise _err("format", "symbol")
    close = _num(data.get("close"))
    if close is None:
        raise _err("missing_field", "close")
    if close <= 0:
        raise _err("format", "close")
    return Quote(
        close,
        _num_or_none(data.get("price_change")),
        _num_or_none(data.get("price_change_percent")),
        _aware(data.get("last_trade_time"), _ny(), "last_trade_time"),
        _aware(root.get("timestamp"), UTC, "timestamp"),
        symbol,
    )


# ── 美債、CFTC ───────────────────────────────────────────────────────────
_TREASURY_COLUMNS = (("m3", "3 mo"), ("y2", "2 yr"), ("y10", "10 yr"), ("y30", "30 yr"))


def parse_treasury_csv(text: object) -> list[YieldRow]:
    """財政部日殖利率 CSV：Date,"1 Mo",…,"3 Mo",…,"2 Yr",…,"10 Yr",…,"30 Yr"（新日期在上）；由舊到新。
    債市休市的日子（例：哥倫布日）本來就沒有列，日變化用這個序列自己的前一列（PLAN §18.8）。
    整份空白 → empty；只有表頭 → []。"""
    body = _as_text(text)
    try:
        lines = list(csv.reader(io.StringIO(body)))
    except csv.Error:
        raise _err("format", "csv") from None
    lines = [line for line in lines if any(cell.strip() for cell in line)]
    if not lines:
        raise _err("empty", "csv")
    header = [cell.strip().lower() for cell in lines[0]]
    if not header or header[0] != "date":
        raise _err("format", "header")
    columns = {key: header.index(label) for key, label in _TREASURY_COLUMNS if label in header}
    if len(columns) != len(_TREASURY_COLUMNS):
        raise _err("format", "columns")
    found: dict[date, YieldRow] = {}
    for line in lines[1:]:
        if len(line) != len(header):
            raise _err("format", "row")
        values: dict[str, Optional[float]] = {}
        for key, index in columns.items():
            rate = _num(line[index])
            if rate is not None and not -5.0 <= rate <= 50.0:
                raise _err("format", key)
            values[key] = rate
        entry = YieldRow(_mdy(line[0], "Date"), **values)
        if entry.date in found and found[entry.date] != entry:
            raise _err("conflict", "Date")
        found[entry.date] = entry
    return [found[day] for day in sorted(found)]  # 只有表頭（月初還沒公布）→ []，不算失敗


CFTC_COLUMNS = 87
# 兩個合約一定用同一種口徑（PLAN §3.1）：檔裡兩個 Consolidated 都有就用 Consolidated，否則都用 E-mini。
CFTC_CODES: dict[str, dict[str, str]] = {
    "consolidated": {"es": "13874+", "nq": "20974+"},
    "e-mini": {"es": "13874A", "nq": "209742"},
}


def _cftc_row(cells: list[str]) -> Optional[CftcRow]:
    """一列（1 起算）：1 名稱、2 YYMMDD、3 YYYY-MM-DD、4 合約代號、8 OI、9–20 Dealer／AM／Lev／Other 的
    long／short／spread、21–22 合計、23–24 非報告、25–41 變化欄。各類加總要等於合計（防欄位位移）。"""
    try:
        report = date.fromisoformat(cells[2].strip())
    except ValueError:
        return None
    if cells[1].strip() != report.strftime("%y%m%d"):
        return None
    try:
        numbers = [_int(cell.strip() if cell.strip() != "." else None) for cell in cells[7:41]]
    except ParseError:
        return None
    positions = numbers[:17]  # 第 8–24 欄
    if any(value is None or value < 0 for value in positions):
        return None
    oi, dl, ds, dsp, al, as_, asp, ll, ls, lsp, ol, os_, osp, tl, ts, nl, ns = positions
    spreads = dsp + asp + lsp + osp
    if dl + al + ll + ol + spreads != tl or ds + as_ + ls + os_ + spreads != ts or tl + nl != oi:
        return None
    changes = numbers[17:]  # 第 25–41 欄
    market = _norm_text(cells[0], 120)
    return CftcRow(
        market, report, oi, dl, ds, al, as_, ll, ls, ol, os_,
        code=cells[3].strip(),
        nonrept_l=nl,
        nonrept_s=ns,
        chg_oi=changes[0],
        chg_dealer_l=changes[1],
        chg_dealer_s=changes[2],
        chg_am_l=changes[4],
        chg_am_s=changes[5],
        chg_lev_l=changes[7],
        chg_lev_s=changes[8],
        chg_other_l=changes[10],
        chg_other_s=changes[11],
    )


def parse_cftc_tff(text: object) -> list[CftcRow]:
    """CFTC FinFutWk.txt（Traders in Financial Futures，期貨）：沒有表頭、每列 87 欄（引號裡可以有逗號）。
    欄數不對或加總對不上的列略過；整份沒有一列可用 → format。"""
    body = _as_text(text)
    try:
        lines = list(csv.reader(io.StringIO(body)))
    except csv.Error:
        raise _err("format", "csv") from None
    rows: list[CftcRow] = []
    seen = 0
    for cells in lines:
        if not any(cell.strip() for cell in cells):
            continue
        seen += 1
        if len(cells) != CFTC_COLUMNS:
            continue
        row = _cftc_row(cells)
        if row is not None:
            rows.append(row)
    if not seen:
        raise _err("empty", "rows")
    if not rows:
        raise _err("format", "rows")
    keyed: dict[tuple[str, date], CftcRow] = {}
    for row in rows:
        key = (row.code, row.report_date)
        if key in keyed and keyed[key] != row:
            raise _err("conflict", "code")
        keyed[key] = row
    return list(keyed.values())


def pick_cftc_contracts(rows: Iterable[CftcRow]) -> CftcPair:
    """挑 S&P 500 與 Nasdaq-100 兩列（同一種口徑、同一個報告日；各取最新一週）。"""
    latest: dict[str, CftcRow] = {}
    for row in rows:
        if row.code not in latest or row.report_date > latest[row.code].report_date:
            latest[row.code] = row
    for basis in ("consolidated", "e-mini"):
        es = latest.get(CFTC_CODES[basis]["es"])
        nq = latest.get(CFTC_CODES[basis]["nq"])
        if es is not None and nq is not None:
            if es.report_date != nq.report_date:
                raise _err("conflict", "report_date")
            return CftcPair(es, nq, basis)
    raise _err("missing_field", "contracts")


# ── Fed FOMC ─────────────────────────────────────────────────────────────
class _FomcParser(HTMLParser):
    """依序收「年份標題」「月份」「日期」三種片段（行事曆頁的 h4 與 fomc-meeting__month／__date）。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.items: list[tuple[str, str]] = []
        self._depth = 0
        self._capture: Optional[tuple[str, int]] = None
        self._buffer: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        if tag == "div":
            self._depth += 1
            if self._capture is None:
                classes = (dict(attrs).get("class") or "").split()
                if "fomc-meeting__month" in classes:
                    self._start("month", self._depth)
                elif "fomc-meeting__date" in classes:
                    self._start("date", self._depth)
        elif tag in ("h3", "h4", "h5") and self._capture is None:
            self._start("heading", -1)

    def handle_endtag(self, tag: str) -> None:
        if tag == "div":
            if self._capture is not None and self._capture[1] == self._depth:
                self._emit()
            self._depth = max(0, self._depth - 1)
        elif tag in ("h3", "h4", "h5") and self._capture is not None and self._capture[0] == "heading":
            self._emit()

    def handle_data(self, data: str) -> None:
        if self._capture is not None:
            self._buffer.append(data)

    def _start(self, kind: str, depth: int) -> None:
        self._capture = (kind, depth)
        self._buffer = []

    def _emit(self) -> None:
        if self._capture is not None:
            self.items.append((self._capture[0], " ".join("".join(self._buffer).split())))
        self._capture = None
        self._buffer = []


_FOMC_YEAR_RE = re.compile(r"(\d{4}) FOMC Meetings")
_FOMC_DAYS_RE = re.compile(r"(\d{1,2})(?:\s*[-–]\s*(\d{1,2}))?")


def _fomc_meeting(year: int, month_text: str, date_text: str) -> Optional[FomcMeeting]:
    """「January」「27-28」、「Apr/May」「30-1」、「March」「17-18*」；「(notation vote)」這類不是會議。"""
    months = [_MONTHS.get(part.strip()[:3].lower()) for part in month_text.split("/")]
    if not months or len(months) > 2 or any(month is None for month in months):
        return None
    clean = date_text.replace("*", "").strip()
    match = _FOMC_DAYS_RE.fullmatch(clean)
    if not match:
        return None
    first, last = int(match.group(1)), int(match.group(2) or match.group(1))
    start_month, end_month = months[0], months[-1]
    end_year = year + 1 if end_month < start_month else year
    try:
        start, end = date(year, start_month, first), date(end_year, end_month, last)
    except ValueError:
        return None
    if not timedelta(0) <= end - start <= timedelta(days=3):
        return None
    return FomcMeeting(start, end, "*" in date_text)


def parse_fed_fomc(html_text: object, *, min_meetings: int = 4) -> list[FomcMeeting]:
    """federalreserve.gov/monetarypolicy/fomccalendars.htm：每年一個面板（2026、2025…，未來年份 2027
    排在最後）。回傳全部會期，依開始日排序。"""
    text = _as_text(html_text)
    parser = _FomcParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception:
        raise _err("format", "html") from None
    year: Optional[int] = None
    month: Optional[str] = None
    meetings: set[FomcMeeting] = set()
    for kind, value in parser.items:
        if kind == "heading":
            match = _FOMC_YEAR_RE.fullmatch(value)
            if match:
                year, month = int(match.group(1)), None
        elif year is None:
            continue
        elif kind == "month":
            month = value
        elif kind == "date":
            meeting = _fomc_meeting(year, month, value) if month else None
            month = None
            if meeting is not None:
                meetings.add(meeting)
    if not meetings:
        raise _err("format", "meetings")
    if len(meetings) < min_meetings:
        raise _err("too_few", "meetings")
    return sorted(meetings, key=lambda m: (m.start, m.end))


# ── 抓取 ──────────────────────────────────────────────────────────────────
def _transport(http: Any) -> Any:
    """抓取層：預設 market_http（用到才 import，這個模組沒有它也能 import 與測試）。"""
    if http is not None:
        return http
    import market_http  # noqa: PLC0415

    return market_http


def _fetch(key: str, url: str, http: Any, ledger: Any) -> bytes:
    endpoint = ENDPOINTS[key]
    response = _transport(http).fetch(
        url,
        max_bytes=endpoint.max_bytes,
        timeout=endpoint.timeout,
        accept=endpoint.accept,
        ledger=ledger,
    )
    body = getattr(response, "body", None)
    if not isinstance(body, (bytes, bytearray)):
        raise _err("format", "body")
    if len(body) > endpoint.max_bytes:
        raise _err("format", "body_size")  # 傳輸層已限制；假的或換掉的傳輸層也不能塞超大內容
    return bytes(body)


def _nasdaq_url(path: str, params: Optional[dict[str, Any]] = None) -> str:
    return f"{NASDAQ_API}/{path}" + (f"?{urlencode(params)}" if params else "")


def fetch_nasdaq_historical(
    symbol: str,
    start: date,
    end: date,
    *,
    etf: bool = False,
    http: Any = None,
    ledger: Any = None,
) -> list[Bar]:
    """一檔的日線（ETF 用 assetclass=etf）。只回 [start, end] 內的列：job 傳 end＝D，晚於 D 的列
    （盤中、隔天亞洲時段）一律不收（PLAN §18.8）。"""
    if start > end:
        raise ValueError("start after end")
    path_symbol = _path_symbol(symbol)
    params = {
        "assetclass": "etf" if etf else "stocks",
        "fromdate": start.isoformat(),
        "todate": end.isoformat(),
        "limit": min((end - start).days + 1, 9999),
    }
    body = _fetch("nasdaq_historical", _nasdaq_url(f"quote/{path_symbol}/historical", params), http, ledger)
    bars = parse_nasdaq_historical(body, expect_symbol=symbol)
    return [bar for bar in bars if start <= bar.date <= end]


def fetch_nasdaq_market_info(*, http: Any = None, ledger: Any = None) -> MarketInfo:
    return parse_nasdaq_market_info(_fetch("nasdaq_market_info", _nasdaq_url("market-info"), http, ledger))


def fetch_nasdaq_splits(since: date, *, http: Any = None, ledger: Any = None) -> list[Split]:
    """executionDate ≥ since 的分割事件；bootstrap 用歷史起點、之後每天 D−7（coverage 由 job 記）。"""
    url = _nasdaq_url("calendar/splits", {"date": since.isoformat()})
    return parse_nasdaq_splits(_fetch("nasdaq_splits", url, http, ledger))


def fetch_nasdaq100(*, http: Any = None, ledger: Any = None) -> HoldingList:
    body = _fetch("nasdaq_nasdaq100", _nasdaq_url("quote/list-type/nasdaq100"), http, ledger)
    return parse_nasdaq100(body)


def fetch_nasdaq_13f(symbol: str, *, http: Any = None, ledger: Any = None) -> Inst13F:
    params = {"limit": 10, "type": "TOTAL", "sortColumn": "marketValue"}
    url = _nasdaq_url(f"company/{_path_symbol(symbol)}/institutional-holdings", params)
    return parse_nasdaq_13f(_fetch("nasdaq_13f", url, http, ledger))


def fetch_nasdaq_financials(symbol: str, *, http: Any = None, ledger: Any = None) -> list[Quarter]:
    url = _nasdaq_url(f"company/{_path_symbol(symbol)}/financials", {"frequency": 2})
    return parse_nasdaq_financials(_fetch("nasdaq_financials", url, http, ledger))


def fetch_nasdaq_summary_market_cap(symbol: str, *, http: Any = None, ledger: Any = None) -> Optional[float]:
    """一檔的市值（美元，沒有回 None）；路徑代號照 canonical（BRK.B 可以直接查，實測 10/18）。"""
    url = _nasdaq_url(f"quote/{_path_symbol(symbol)}/summary", {"assetclass": "stocks"})
    return parse_nasdaq_summary_market_cap(_fetch("nasdaq_summary", url, http, ledger), expect_symbol=symbol)


def fetch_nasdaq_econ_calendar(
    day: date, *, http: Any = None, ledger: Any = None, require_sentinel: bool = False
) -> list[EconEvent]:
    """美東 `day` 的經濟事件。Nasdaq 的 `date=X` 回的是 X 前一天的事件（2026-10-18 實測：date=10/21 回週二的
    Redbook、API 原油庫存，date=10/22 回週三的 MBA 房貸、EIA 原油庫存，date=10/23 回週四的初領失業金、
    Fed 資產負債表），所以要 day 的事件就請求 date=day＋1；解析器再用週三的 MBA 驗一次。"""
    url = _nasdaq_url("calendar/economicevents", {"date": (day + timedelta(days=1)).isoformat()})
    return parse_nasdaq_econ_calendar(_fetch("nasdaq_econ_calendar", url, http, ledger), day, require_sentinel=require_sentinel)


def fetch_nasdaq_earnings_date(symbol: str, *, http: Any = None, ledger: Any = None) -> EarningsDate:
    url = _nasdaq_url(f"analyst/{_path_symbol(symbol)}/earnings-date")
    return parse_nasdaq_earnings_date(_fetch("nasdaq_earnings", url, http, ledger), expect_symbol=symbol)


def fetch_ssga_spy(*, http: Any = None, ledger: Any = None) -> HoldingList:
    return parse_ssga_spy_xlsx(_fetch("ssga_spy", SSGA_SPY_URL, http, ledger))


def fetch_moneydj_holdings(etf: str, *, http: Any = None, ledger: Any = None) -> HoldingList:
    etfid, _ = _moneydj_etf(etf)
    url = f"{MONEYDJ_URL}?{urlencode({'etfid': etfid})}"
    return parse_moneydj_holdings(_fetch("moneydj", url, http, ledger), etfid)


def fetch_cboe_options(symbol: str, *, http: Any = None, ledger: Any = None) -> OptionChain:
    if symbol not in CBOE_OPTION_SYMBOLS:
        raise _err("format", "symbol")
    url = CBOE_OPTIONS_URL.format(symbol=quote(symbol, safe=""))
    return parse_cboe_options(_fetch("cboe_options", url, http, ledger), expect_symbol=symbol)


def fetch_cboe_quote(symbol: str, *, http: Any = None, ledger: Any = None) -> Quote:
    if symbol not in CBOE_QUOTE_SYMBOLS:
        raise _err("format", "symbol")
    url = CBOE_QUOTE_URL.format(symbol=quote(symbol, safe=""))
    return parse_cboe_quote(_fetch("cboe_quote", url, http, ledger), expect_symbol=symbol)


def fetch_treasury(month: date, *, http: Any = None, ledger: Any = None) -> list[YieldRow]:
    """某個月（取 month 的年月）的日殖利率；月初還沒有列時回 []，job 另抓上個月才有前一列算日變化。"""
    if not isinstance(month, date) or not 1990 <= month.year <= 2200:
        raise ValueError("bad month")
    url = TREASURY_URL.format(ym=f"{month.year:04d}{month.month:02d}")
    rows = parse_treasury_csv(_fetch("treasury", url, http, ledger))
    return [row for row in rows if (row.date.year, row.date.month) == (month.year, month.month)]


def fetch_cftc_tff(*, http: Any = None, ledger: Any = None) -> list[CftcRow]:
    return parse_cftc_tff(_fetch("cftc_tff", CFTC_TFF_URL, http, ledger))


def fetch_finra_margin(*, http: Any = None, ledger: Any = None) -> list[MarginRow]:
    return parse_finra_margin_xlsx(_fetch("finra_margin", FINRA_MARGIN_URL, http, ledger))


def fetch_fed_fomc(*, http: Any = None, ledger: Any = None) -> list[FomcMeeting]:
    return parse_fed_fomc(_fetch("fed_fomc", FED_FOMC_URL, http, ledger))
